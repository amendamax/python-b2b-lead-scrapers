import sys
import os
import re
import time
import random
import argparse
import base64
import urllib.parse
from curl_cffi import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Force UTF-8 encoding on Windows console to avoid encoding crashes
try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except AttributeError:
    pass

class B2BAgencyOutreachScraper:
    def __init__(self):
        self.session = requests.Session(impersonate="chrome110")
        self.headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept-Language": "it-IT,it;q=0.9,en-US;q=0.8,en;q=0.7",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8",
            "Referer": "https://www.google.it/"
        }
        self.ignore_domains = [
            "facebook.com", "instagram.com", "linkedin.com", "youtube.com", "twitter.com", "pinterest.com", 
            "wikipedia.org", "google.com", "duckduckgo.com", "bing.com", "yahoo.com", "yandex.com",
            "immobiliare.it", "idealista.it", "casa.it", "tecnocasa.it", "tempocasa.it", "gabetti.it",
            "paginegialle.it", "yelp.com", "yellowpages.com", "tripadvisor.com"
        ]

    def is_ignored(self, domain):
        domain_low = domain.lower()
        return any(x in domain_low for x in self.ignore_domains)

    def clean_domain(self, url):
        try:
            parsed = urllib.parse.urlparse(url)
            domain = parsed.netloc.replace("www.", "")
            return domain
        except:
            return ""

    def search_duckduckgo(self, query, target_count):
        print(f"  [Search] Căutare pe DuckDuckGo pentru: {query}")
        links = []
        try:
            encoded_query = urllib.parse.quote_plus(query)
            url = f"https://html.duckduckgo.com/html/?q={encoded_query}"
            r = self.session.get(url, headers=self.headers, timeout=12)
            if r.status_code != 200:
                return []
                
            soup = BeautifulSoup(r.text, 'html.parser')
            for a in soup.select('a.result__url'):
                href = a.get('href', '')
                if 'uddg=' in href:
                    actual_url = href.split('uddg=')[1].split('&')[0]
                    actual_url = urllib.parse.unquote(actual_url)
                    if actual_url.startswith('http'):
                        links.append(actual_url)
        except Exception as e:
            pass
        return links

    def clean_leads_candidates(self, links, city, country):
        candidates = []
        seen_domains = set()
        for link in links:
            domain = self.clean_domain(link)
            if not domain or self.is_ignored(domain):
                continue
            if domain not in seen_domains:
                seen_domains.add(domain)
                name_parts = domain.split(".")[0].replace("-", " ").replace("_", " ").title()
                candidates.append({
                    "name": name_parts,
                    "website": f"https://{domain}",
                    "location": f"{city}, {country}"
                })
        return candidates

    def crawl_agency_website(self, company_name, url):
        print(f"    [Crawler] Scanăm site-ul: {url}...")
        domain = self.clean_domain(url)
        emails = []
        
        try:
            r = self.session.get(url, headers=self.headers, timeout=10)
            if r.status_code not in [200, 301, 302]:
                if url.startswith('https://'):
                    url_http = url.replace('https://', 'http://')
                    r = self.session.get(url_http, headers=self.headers, timeout=8)
            html = r.text
            emails.extend(re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', html))
        except:
            pass
            
        valid_emails = []
        for email in emails:
            email_low = email.strip().lower()
            if not re.match(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$', email_low):
                continue
            if any(email_low.endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".gif", ".svg"]):
                continue
            if any(x in email_low for x in ["example", "domain.com", "yourdomain", "sentry"]):
                continue
            valid_emails.append(email_low)
            
        unique_emails = list(set(valid_emails))
        def email_priority(e):
            if domain in e:
                if any(handle in e for handle in ["info@", "contatt", "office@"]):
                    return 0
                return 1
            return 2
        unique_emails.sort(key=email_priority)
        
        final_email = unique_emails[0] if unique_emails else "N/A"
        return final_email

    def run_pipeline(self, target_count):
        print("\n=======================================================")
        print("🚀 EXTRAGERE AGENȚII IMOBILIARE ITALIA (TEST)")
        print("=======================================================\n")
        
        targets = [
            {"city": "Firenze", "country": "Italy", "query": '"agenzia immobiliare" Firenze'},
            {"city": "Bologna", "country": "Italy", "query": '"agenzia immobiliare" Bologna'},
            {"city": "Napoli", "country": "Italy", "query": '"agenzia immobiliare" Napoli'},
            {"city": "Genova", "country": "Italy", "query": '"agenzia immobiliare" Genova'},
            {"city": "Verona", "country": "Italy", "query": '"agenzia immobiliare" Verona'}
        ]
        
        leads = []
        
        for tgt in targets:
            if len(leads) >= target_count:
                break
            city = tgt["city"]
            query = tgt["query"]
            
            links = self.search_duckduckgo(query, target_count)
            candidates = self.clean_leads_candidates(links, city, "Italy")
            
            for cand in candidates:
                if len(leads) >= target_count:
                    break
                    
                name = cand["name"]
                web = cand["website"]
                loc = cand["location"]
                
                email = self.crawl_agency_website(name, web)
                if email == "N/A":
                    continue
                    
                leads.append({
                    "company_name": name,
                    "location": loc,
                    "website": web,
                    "email": email
                })
                print(f"     [CALIFICAT] {name} | Email: {email}")
                
        return leads

    def format_excel_report(self, leads, output_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Agentii Imobiliare Italia"
        
        headers = ["Nume Agenție", "Locație", "Site Web", "E-mail Contact"]
        for col_idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=h)
            
        for row_idx, lead in enumerate(leads, start=2):
            ws.cell(row=row_idx, column=1, value=lead["company_name"])
            ws.cell(row=row_idx, column=2, value=lead["location"])
            ws.cell(row=row_idx, column=3, value=lead["website"])
            ws.cell(row=row_idx, column=4, value=lead["email"])
            
        for col in ws.columns:
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = 30
            
        wb.save(output_path)

def main():
    scraper = B2BAgencyOutreachScraper()
    leads = scraper.run_pipeline(14)
    
    if leads:
        filename = "Agenzie_Immobiliari_Italia_Part2.xlsx"
        desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        desktop_path = os.path.join(desktop_dir, filename)
        scraper.format_excel_report(leads, desktop_path)
        print(f"\n🚀 RAPORT FINALIZAT SALVAT PE DESKTOP: {desktop_path}")
    else:
        print("Nu s-au găsit date.")

if __name__ == "__main__":
    main()
